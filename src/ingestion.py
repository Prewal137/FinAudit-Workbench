"""Excel workbook ingestion engine for FinAudit-Workbench.

Implements dual-pass loading with openpyxl, explicit row-by-row mapping,
period normalization, and cell value provenance as specified in
docs/workbook_ingestion_spec.md.
"""

from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional, Tuple, Union

import openpyxl
from openpyxl.utils import get_column_letter

from src.models import (
    CellValue,
    FinancialLineItem,
    FinancialStatement,
    PeriodHeader,
    RowClassification,
)


# Expected sheet dimensions
EXPECTED_IS_MAX_ROW = 107
EXPECTED_IS_MAX_COL = 25
EXPECTED_BS_MAX_ROW = 71
EXPECTED_BS_MAX_COL = 26

# Explicit row classification maps
IS_ROW_CLASSIFICATION: Dict[int, RowClassification] = {
    # Metadata & structural
    1: RowClassification.METADATA,
    2: RowClassification.METADATA,
    3: RowClassification.METADATA,
    4: RowClassification.SPACER,
    5: RowClassification.HEADER,
    6: RowClassification.SPACER,
    # Income partition
    7: RowClassification.SECTION_HEADING,
    **{r: RowClassification.DETAIL_ACCOUNT for r in range(8, 14)},
    14: RowClassification.SUBTOTAL,
    15: RowClassification.SPACER,
    # COGS partition
    16: RowClassification.SECTION_HEADING,
    **{r: RowClassification.DETAIL_ACCOUNT for r in range(17, 24)},
    24: RowClassification.SUBTOTAL,
    25: RowClassification.SPACER,
    # Gross Profit
    26: RowClassification.CALCULATED_RESULT,
    27: RowClassification.SPACER,
    # OpEx partition
    28: RowClassification.SECTION_HEADING,
    **{r: RowClassification.DETAIL_ACCOUNT for r in range(29, 47)},
    47: RowClassification.SUBSECTION_HEADING,
    **{r: RowClassification.DETAIL_ACCOUNT for r in range(48, 59)},
    59: RowClassification.SUBTOTAL,
    60: RowClassification.SUBSECTION_HEADING,
    **{r: RowClassification.DETAIL_ACCOUNT for r in range(61, 71)},
    71: RowClassification.SUBTOTAL,
    **{r: RowClassification.DETAIL_ACCOUNT for r in range(72, 103)},
    103: RowClassification.SUBTOTAL,
    104: RowClassification.SPACER,
    # Operating & Net Results
    105: RowClassification.CALCULATED_RESULT,
    106: RowClassification.SPACER,
    107: RowClassification.CALCULATED_RESULT,
}

BS_ROW_CLASSIFICATION: Dict[int, RowClassification] = {
    # Metadata & structural
    1: RowClassification.METADATA,
    2: RowClassification.METADATA,
    3: RowClassification.METADATA,
    4: RowClassification.SPACER,
    5: RowClassification.HEADER,
    6: RowClassification.SPACER,
    # Assets partition
    7: RowClassification.SECTION_HEADING,
    8: RowClassification.SUBSECTION_HEADING,
    9: RowClassification.SUBSECTION_HEADING,
    **{r: RowClassification.DETAIL_ACCOUNT for r in range(10, 15)},
    15: RowClassification.SUBTOTAL,
    **{r: RowClassification.DETAIL_ACCOUNT for r in range(16, 27)},
    27: RowClassification.SUBTOTAL,
    28: RowClassification.SUBSECTION_HEADING,
    **{r: RowClassification.DETAIL_ACCOUNT for r in range(29, 33)},
    33: RowClassification.SUBTOTAL,
    34: RowClassification.SUBSECTION_HEADING,
    35: RowClassification.DETAIL_ACCOUNT,
    36: RowClassification.CALCULATED_RESULT,
    37: RowClassification.CALCULATED_RESULT,
    38: RowClassification.SPACER,
    # Liabilities & Equity partition
    39: RowClassification.SECTION_HEADING,
    40: RowClassification.SUBSECTION_HEADING,
    41: RowClassification.SUBSECTION_HEADING,
    **{r: RowClassification.DETAIL_ACCOUNT for r in range(42, 54)},
    54: RowClassification.SUBTOTAL,
    55: RowClassification.SUBSECTION_HEADING,
    56: RowClassification.DETAIL_ACCOUNT,
    57: RowClassification.DETAIL_ACCOUNT,
    58: RowClassification.SUBTOTAL,
    59: RowClassification.CALCULATED_RESULT,
    60: RowClassification.SUBSECTION_HEADING,
    **{r: RowClassification.DETAIL_ACCOUNT for r in range(61, 70)},
    70: RowClassification.SUBTOTAL,
    71: RowClassification.CALCULATED_RESULT,
}

IS_DORMANT_ROWS = {89, 90, 91, 92}


def normalize_period_header(raw_header: str) -> str:
    """Normalize raw period strings (e.g., 'Jun 2026', '30 Jun 2026') to 'YYYY-MM'.
    
    Raises ValueError if header cannot be parsed into standard date.
    """
    cleaned = raw_header.strip()
    # Try '30 Jun 2026' pattern
    for fmt in ("%d %b %Y", "%b %Y", "%B %Y", "%Y-%m", "%m/%d/%Y"):
        try:
            dt = datetime.strptime(cleaned, fmt)
            return dt.strftime("%Y-%m")
        except ValueError:
            continue
    raise ValueError(f"Unable to parse period header: {raw_header!r}")


def _determine_hierarchy_path_is(row_num: int) -> Tuple[str, ...]:
    """Map Income Statement row to its category hierarchy path."""
    if 7 <= row_num <= 14:
        return ("INCOME",)
    if 16 <= row_num <= 24:
        return ("COST OF GOODS SOLD",)
    if row_num == 26:
        return ()
    if 28 <= row_num <= 46:
        return ("EXPENSES", "Direct OpEx")
    if 47 <= row_num <= 59:
        return ("EXPENSES", "Sales & Marketing Payroll")
    if 60 <= row_num <= 71:
        return ("EXPENSES", "Marketing")
    if 72 <= row_num <= 103:
        return ("EXPENSES", "Other OpEx")
    if row_num in (105, 107):
        return ()
    return ()


def _determine_hierarchy_path_bs(row_num: int) -> Tuple[str, ...]:
    """Map Balance Sheet row to its category hierarchy path."""
    if row_num in range(7, 38):
        # Assets partition
        if row_num in range(8, 28):
            if row_num in range(9, 16):
                return ("ASSETS", "Current Assets", "Cash and Cash Equivalents")
            return ("ASSETS", "Current Assets")
        if row_num in range(28, 34):
            return ("ASSETS", "Fixed Assets")
        if row_num in range(34, 37):
            return ("ASSETS", "Long Term Assets")
        return ("ASSETS",)
    if row_num in range(39, 72):
        # Liabilities & Equity partition
        if row_num in range(40, 55):
            if row_num in range(41, 54):
                return ("LIABILITIES & EQUITY", "Current Liabilities", "Other Current Liabilities")
            return ("LIABILITIES & EQUITY", "Current Liabilities")
        if row_num in range(55, 59):
            return ("LIABILITIES & EQUITY", "Long Term Liabilities")
        if row_num == 59:
            return ("LIABILITIES & EQUITY", "Total Liabilities")
        if row_num in range(60, 71):
            return ("LIABILITIES & EQUITY", "Equity")
        return ("LIABILITIES & EQUITY",)
    return ()


def _extract_row_label(ws, row_num: int, sheet_name: str) -> Tuple[str, str]:
    """Extract original exact label and normalized label.
    
    Checks columns A, B, and C in order of hierarchy for the sheet.
    """
    val_a = ws.cell(row_num, 1).value
    val_b = ws.cell(row_num, 2).value
    val_c = ws.cell(row_num, 3).value

    # Spacers have no label
    classification = (
        IS_ROW_CLASSIFICATION.get(row_num)
        if sheet_name == "Income Statement"
        else BS_ROW_CLASSIFICATION.get(row_num)
    )
    if classification == RowClassification.SPACER:
        return ("", "")

    # For IS: metadata and sections in Col A; detail, subtotals, subsections in Col B; calc in Col A
    if sheet_name == "Income Statement":
        if row_num in (1, 2, 3, 7, 16, 26, 28, 105, 107):
            raw = str(val_a) if val_a is not None else ""
        else:
            raw = str(val_b) if val_b is not None else (str(val_a) if val_a is not None else "")
    else:
        # Balance Sheet
        if row_num in (1, 2, 3, 7, 37, 39, 71):
            raw = str(val_a) if val_a is not None else ""
        elif row_num in (8, 27, 28, 33, 34, 36, 40, 54, 58, 59, 60, 70):
            raw = str(val_b) if val_b is not None else (str(val_a) if val_a is not None else "")
        else:
            # Accounts, cash subtotal, subsections in Col C
            raw = str(val_c) if val_c is not None else (str(val_b) if val_b is not None else "")

    original_label = raw
    normalized_label = raw.strip()
    return (original_label, normalized_label)


def ingest_statement(
    ws_formula,
    ws_values,
    sheet_name: str,
    period_start_col: int,
    period_end_col: int,
    trailing_start_col: int,
    trailing_end_col: int,
    row_classification_map: Dict[int, RowClassification],
    max_expected_row: int,
) -> FinancialStatement:
    """Ingest a single financial worksheet into a strongly-typed FinancialStatement."""
    # 1. Verify period headers in Row 5
    periods: List[PeriodHeader] = []
    for col in range(period_start_col, period_end_col + 1):
        cell_val = ws_values.cell(5, col).value
        if cell_val is None:
            raise ValueError(f"Missing period header in {sheet_name} at cell {get_column_letter(col)}5")
        raw_header = str(cell_val)
        norm_period = normalize_period_header(raw_header)
        periods.append(
            PeriodHeader(
                original_header=raw_header,
                normalized_period=norm_period,
                column_index=col,
                column_letter=get_column_letter(col),
            )
        )

    # 2. Check trailing columns have zero populated cells
    trailing_populated = []
    for r in range(1, max_expected_row + 1):
        for c in range(trailing_start_col, trailing_end_col + 1):
            val_f = ws_formula.cell(r, c).value
            val_v = ws_values.cell(r, c).value
            if val_f is not None or val_v is not None:
                trailing_populated.append(f"{get_column_letter(c)}{r}")
    
    trailing_empty_ok = (len(trailing_populated) == 0)

    # 3. Ingest row-by-row
    items: List[FinancialLineItem] = []
    for row_num in range(1, max_expected_row + 1):
        classification = row_classification_map.get(row_num, RowClassification.DETAIL_ACCOUNT)
        orig_label, norm_label = _extract_row_label(ws_formula, row_num, sheet_name)
        
        hierarchy = (
            _determine_hierarchy_path_is(row_num)
            if sheet_name == "Income Statement"
            else _determine_hierarchy_path_bs(row_num)
        )

        is_dormant = (sheet_name == "Income Statement" and row_num in IS_DORMANT_ROWS)

        # Ingest values across all period columns
        values_by_period: Dict[str, CellValue] = {}
        for period_header in periods:
            col_idx = period_header.column_index
            c_f = ws_formula.cell(row_num, col_idx)
            c_v = ws_values.cell(row_num, col_idx)

            raw_f = c_f.value
            raw_v = c_v.value

            if isinstance(raw_f, str) and raw_f.startswith("="):
                formula_text = raw_f
                literal_val = None
                cached_val = raw_v
            else:
                formula_text = None
                literal_val = raw_f
                cached_val = raw_v

            values_by_period[period_header.normalized_period] = CellValue(
                literal_value=literal_val,
                formula_text=formula_text,
                cached_value=cached_val,
            )

        items.append(
            FinancialLineItem(
                sheet_name=sheet_name,
                row_number=row_num,
                original_label=orig_label,
                normalized_label=norm_label,
                classification=classification,
                hierarchy_path=hierarchy,
                is_dormant=is_dormant,
                values_by_period=values_by_period,
            )
        )

    return FinancialStatement(
        statement_name=sheet_name,
        periods=periods,
        items=items,
        trailing_empty_columns_verified=trailing_empty_ok,
    )


def load_financial_workbook(file_path: Union[str, Path]) -> Tuple[FinancialStatement, FinancialStatement]:
    """Ingest both Income Statement and Balance Sheet from Excel workbook.
    
    Returns:
        (income_statement, balance_sheet)
    """
    path = Path(file_path).resolve()
    if not path.exists():
        raise FileNotFoundError(f"Workbook file not found at {path}")

    # Dual openpyxl loading
    wb_formula = openpyxl.load_workbook(path, data_only=False)
    wb_values = openpyxl.load_workbook(path, data_only=True)

    required_sheets = ["Income Statement", "Balance Sheet"]
    for s in required_sheets:
        if s not in wb_formula.sheetnames:
            raise ValueError(f"Required sheet {s!r} missing from workbook. Present sheets: {wb_formula.sheetnames}")

    ws_is_f, ws_is_v = wb_formula["Income Statement"], wb_values["Income Statement"]
    ws_bs_f, ws_bs_v = wb_formula["Balance Sheet"], wb_values["Balance Sheet"]

    # Verify structural limits
    if ws_is_f.max_row != EXPECTED_IS_MAX_ROW or ws_is_f.max_column != EXPECTED_IS_MAX_COL:
        raise ValueError(
            f"Income Statement dimension mismatch: expected {EXPECTED_IS_MAX_ROW}x{EXPECTED_IS_MAX_COL}, "
            f"got {ws_is_f.max_row}x{ws_is_f.max_column}"
        )
    if ws_bs_f.max_row != EXPECTED_BS_MAX_ROW or ws_bs_f.max_column != EXPECTED_BS_MAX_COL:
        raise ValueError(
            f"Balance Sheet dimension mismatch: expected {EXPECTED_BS_MAX_ROW}x{EXPECTED_BS_MAX_COL}, "
            f"got {ws_bs_f.max_row}x{ws_bs_f.max_column}"
        )

    # Ingest Income Statement: Period Cols C-T (3..20), Trailing Cols U-Y (21..25)
    is_stmt = ingest_statement(
        ws_formula=ws_is_f,
        ws_values=ws_is_v,
        sheet_name="Income Statement",
        period_start_col=3,
        period_end_col=20,
        trailing_start_col=21,
        trailing_end_col=25,
        row_classification_map=IS_ROW_CLASSIFICATION,
        max_expected_row=EXPECTED_IS_MAX_ROW,
    )

    # Ingest Balance Sheet: Period Cols D-U (4..21), Trailing Cols V-Z (22..26)
    bs_stmt = ingest_statement(
        ws_formula=ws_bs_f,
        ws_values=ws_bs_v,
        sheet_name="Balance Sheet",
        period_start_col=4,
        period_end_col=21,
        trailing_start_col=22,
        trailing_end_col=26,
        row_classification_map=BS_ROW_CLASSIFICATION,
        max_expected_row=EXPECTED_BS_MAX_ROW,
    )

    return is_stmt, bs_stmt
