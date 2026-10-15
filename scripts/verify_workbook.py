from pathlib import Path
import sys
import openpyxl
sys.stdout.reconfigure(encoding='utf-8')

ROOT = Path(__file__).resolve().parents[1]
FILE = ROOT / "data" / "input" / "ABC Inc - Financial Statements.xlsx"
TOLERANCE = 0.01

if not FILE.exists():
    raise FileNotFoundError(f"Workbook not found: {FILE}")

# Load formulas and cached results separately.
wb_formula = openpyxl.load_workbook(FILE, data_only=False)
wb_values = openpyxl.load_workbook(FILE, data_only=True)

isf, isv = wb_formula["Income Statement"], wb_values["Income Statement"]
bsf, bsv = wb_formula["Balance Sheet"], wb_values["Balance Sheet"]

print("WORKBOOK:", FILE)
print("\n1. SHEET DIMENSIONS")
for name in wb_formula.sheetnames:
    ws = wb_formula[name]
    print(f"{name}: {ws.max_row} rows x {ws.max_column} columns")

def check(label, actual, expected, cell):
    if not isinstance(actual, (int, float)) or not isinstance(expected, (int, float)):
        print(f"FAIL | {label} | {cell} | non-numeric value")
        return False
    diff = abs(actual - expected)
    status = "PASS" if diff < TOLERANCE else "FAIL"
    print(f"{status} | {label} | {cell} | difference={diff:.10f}")
    return status == "PASS"

print("\n2. MONTHLY HEADERS")
is_periods = [isf.cell(5, c).value for c in range(3, 21)]
bs_periods = [bsf.cell(5, c).value for c in range(4, 22)]
print("Income Statement:", is_periods)
print("Balance Sheet:", bs_periods)
print("Period counts:", len(is_periods), len(bs_periods))
print("Matching periods:", [
    str(a)[:3] == str(b).split()[-2][:3]
    and str(a).split()[-1] == str(b).split()[-1]
    for a, b in zip(is_periods, bs_periods)
])

print("\n3. FORMULAS AND CACHED VALUES")
for name, wsf, wsv in [
    ("Income Statement", isf, isv),
    ("Balance Sheet", bsf, bsv),
]:
    formulas = [
        cell.coordinate
        for row in wsf.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=")
    ]
    missing = [c for c in formulas if wsv[c].value is None]
    print(
        f"{name}: {len(formulas)} formulas; "
        f"{len(missing)} missing cached values"
    )
    if missing:
        print("Missing cache cells:", missing)

print("\n4. INCOME STATEMENT CHECKS")
for c in range(3, 21):
    col = isf.cell(5, c).column_letter
    check("Gross profit", isv[f"{col}26"].value,
          isv[f"{col}14"].value - isv[f"{col}24"].value, f"{col}26")
    check("Operating income", isv[f"{col}105"].value,
          isv[f"{col}26"].value - isv[f"{col}103"].value, f"{col}105")
    check("Net income", isv[f"{col}107"].value,
          isv[f"{col}105"].value, f"{col}107")

print("\n5. BALANCE SHEET CHECKS")
for c in range(4, 22):
    col = bsf.cell(5, c).column_letter
    check("Current assets", bsv[f"{col}27"].value,
          sum(bsv.cell(r, c).value or 0 for r in range(16, 27))
          + bsv[f"{col}15"].value, f"{col}27")
    check("Total assets", bsv[f"{col}37"].value,
          bsv[f"{col}27"].value + bsv[f"{col}33"].value
          + bsv[f"{col}36"].value, f"{col}37")
    check("Total liabilities", bsv[f"{col}59"].value,
          bsv[f"{col}54"].value + bsv[f"{col}58"].value, f"{col}59")
    check("Liabilities and equity", bsv[f"{col}71"].value,
          bsv[f"{col}59"].value + bsv[f"{col}70"].value, f"{col}71")
    check("Assets reconciliation", bsv[f"{col}37"].value,
          bsv[f"{col}71"].value, f"{col}37 vs {col}71")

print("\n6. IMPORTANT SOURCE CELLS")
for cell in ["T8", "T14", "T24", "T26", "I99", "I101", "I103", "I105"]:
    print(f"Income Statement!{cell}: "
          f"formula={isf[cell].value!r}, cached={isv[cell].value!r}")

for cell in ["C35", "C53", "C68"]:
    print(f"Balance Sheet!{cell}: {bsf[cell].value!r}")

print("\n7. TRAILING COLUMN CONTENT")
for name, ws, columns in [
    ("Income Statement", isv, range(21, 26)),
    ("Balance Sheet", bsv, range(22, 27)),
]:
    populated = [
        ws.cell(r, c).coordinate
        for r in range(1, ws.max_row + 1)
        for c in columns
        if ws.cell(r, c).value is not None
    ]
    print(f"{name}: {len(populated)} populated trailing cells")
    if populated:
        print(populated)

print("\nVerification run complete. Review every FAIL and warning.")